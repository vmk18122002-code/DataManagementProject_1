"""Regenerate the SIMULATED Distribution_Records sheet (these are not real deliveries).

The original sheet had 3,000 records, about 8 per DS division, which is too few for an
area's need to show through the randomness. This writes 30,000 records where how often an
area receives an item depends on its real indicators:
  Dry Food            - number of poor people and Need_Score
  Dry Food + Clothes  - number of poor people, how disaster-prone the area is, Need_Score
  Books               - population and district MPI poverty
Random variation is still added per area and per record. The original workbook is copied
to a backup file before the sheet is replaced.
"""
import shutil

import numpy as np
import openpyxl
import pandas as pd

from common import BASE, DIST_FILE, KEY, load_disasters, load_ds_features

SEED = 42
N_RECORDS = 30000
DISASTER_SHARE = 0.12
AREA_NOISE = 0.3                      # random spread in how much each area receives
START, END = "2024-01-02", "2026-10-03"
BACKUP = BASE / "ds_division_distribution_dataset_original_3000_records.xlsx"

ITEM_TYPES = ["Dry Food", "Dry Food + Clothes", "Books"]
NORMAL_MIX = [0.44, 0.31, 0.25]
DISASTER_MIX = [0.51, 0.29, 0.20]
BASE_PACKS = {"Dry Food": 45, "Dry Food + Clothes": 35, "Books": 30}
DISASTER_PACK_FACTOR = 2.2
ORGANIZATIONS = ["World Vision Lanka", "Local Community Trust", "Sri Lanka Red Cross",
                 "Foundation of Goodness", "Rotary Club", "Sarvodaya Shramadana",
                 "District Secretariat Relief Unit", "Lions Club"]
# same disaster-season assumption as the original sheet
DISASTER_DISTRICTS = {"Kandy": 0.26, "Ratnapura": 0.24, "Badulla": 0.22, "Nuwara Eliya": 0.18,
                      "Mannar": 0.06, "Kilinochchi": 0.02, "Mullaitivu": 0.02}
DISASTER_MONTHS = {10: 0.37, 11: 0.18, 12: 0.19, 1: 0.26}

README_DUMMY = [
    "DUMMY data (still synthetic, for project demo only):",
    "  - Distribution_Records sheet: 30,000 simulated donation records (date, item type,",
    "    packs given, families served, mode, organization) are NOT real deliveries.",
    "    Created by generate_distribution_records.py (the first version had 3,000 records,",
    "    kept in ds_division_distribution_dataset_original_3000_records.xlsx).",
    "  - GN Division has been REMOVED as requested - area is now the real DS Division name.",
    "  - How often an area receives each item is simulated from its real indicators:",
    "    Dry Food from poor people + Need_Score; Dry Food + Clothes from poor people +",
    "    past disaster frequency + Need_Score; Books from population + district MPI.",
    "    Random variation is added per area and per record. High-need areas get up to",
    "    ~25% bigger pack quantities. These are modelling choices, not facts, so a model",
    "    trained on this sheet is learning the simulation rule, not real-world demand.",
    "  - Disaster-season clustering (Oct-Jan, flood-prone districts) is the same assumption",
    "    used in the earlier version of this dataset.",
    "  - Item_Type categories: 'Dry Food' (alone), 'Dry Food + Clothes' (combined bundle -",
    "    clothes are never given as a standalone delivery, matching real field reports where",
    "    clothing is usually bundled with a food drop), and 'Books' (alone, for ESSP-style",
    "    school support). Packs_Given for a combo row counts as one bundle per family.",
]


def z(s):
    return (s - s.mean()) / s.std()


def area_weights(ds, rng):
    """Chance of each DS division receiving each item type, from its real indicators."""
    poor, pop = z(np.log(ds["Poor_people"])), z(np.log(ds["Population_2024"]))
    need, mpi = z(ds["Need_Score"]), z(ds["District_MPI_Poor_Pct"])
    prone = z(np.log1p(ds["Disaster_Events_Per_Year"]))
    score = {"Dry Food": 0.9 * poor + 0.5 * need,
             "Dry Food + Clothes": 0.6 * poor + 0.7 * prone + 0.3 * need,
             "Books": 0.9 * pop + 0.5 * mpi}
    out = {}
    for item, s in score.items():
        w = np.exp(s + rng.normal(0, AREA_NOISE, len(ds)))
        out[item] = (w / w.sum()).to_numpy()
    return out


def generate():
    rng = np.random.default_rng(SEED)
    ds = load_ds_features()
    dis = load_disasters()
    years = dis["Year"].max() - dis["Year"].min() + 1
    events = (dis.groupby(KEY).size() / years).rename("Disaster_Events_Per_Year").reset_index()
    ds = ds.merge(events, on=KEY, how="left").fillna({"Disaster_Events_Per_Year": 0})
    weights = area_weights(ds, rng)

    days = pd.date_range(START, END)
    n_dis = int(N_RECORDS * DISASTER_SHARE)
    n_norm = N_RECORDS - n_dis

    # normal time: any day, area chosen by the item's need weights
    item = rng.choice(ITEM_TYPES, n_norm, p=NORMAL_MIX)
    area = np.empty(n_norm, dtype=int)
    for it in ITEM_TYPES:
        m = item == it
        area[m] = rng.choice(len(ds), m.sum(), p=weights[it])
    normal = pd.DataFrame({"Date": rng.choice(days, n_norm), "area": area,
                           "Item_Type": item, "Mode": "Normal"})

    # disaster time: Oct-Jan, in the disaster-prone districts
    day_w = np.array([DISASTER_MONTHS.get(d.month, 0) for d in days])
    dis_district = rng.choice(list(DISASTER_DISTRICTS), n_dis, p=list(DISASTER_DISTRICTS.values()))
    area = np.empty(n_dis, dtype=int)
    for d in DISASTER_DISTRICTS:
        m = dis_district == d
        idx = np.flatnonzero(ds["District"] == d)
        w = (ds["Population_2024"] * (0.5 + ds["Need_Score"])).to_numpy()[idx]
        area[m] = rng.choice(idx, m.sum(), p=w / w.sum())
    disaster = pd.DataFrame({"Date": rng.choice(days, n_dis, p=day_w / day_w.sum()), "area": area,
                             "Item_Type": rng.choice(ITEM_TYPES, n_dis, p=DISASTER_MIX),
                             "Mode": "Disaster"})

    rec = pd.concat([normal, disaster]).sample(frac=1, random_state=SEED).reset_index(drop=True)
    need = ds["Need_Score"].to_numpy()[rec["area"]]
    size = (rec["Item_Type"].map(BASE_PACKS) * np.where(rec["Mode"] == "Disaster", DISASTER_PACK_FACTOR, 1)
            * (1 + 0.25 * need) * rng.lognormal(0, 0.35, len(rec)))
    rec["Packs_Given"] = size.round().clip(lower=5).astype(int)
    rec["Affected_Families"] = (rec["Packs_Given"] * rng.uniform(0.8, 1.0, len(rec))).round().clip(lower=1).astype(int)
    rec["Organization"] = rng.choice(ORGANIZATIONS, len(rec))
    rec["District"] = ds["District"].to_numpy()[rec["area"]]
    rec["DS_Division"] = ds["DS_Division"].to_numpy()[rec["area"]]
    rec["Date"] = pd.to_datetime(rec["Date"]).dt.strftime("%Y-%m-%d")
    rec.insert(0, "No", range(1, len(rec) + 1))
    return rec[["No", "Date", "District", "DS_Division", "Item_Type", "Packs_Given",
                "Affected_Families", "Mode", "Organization"]]


def district_summary(rec):
    food = rec["Item_Type"].isin(["Dry Food", "Dry Food + Clothes"])
    g = lambda mask: rec[mask].groupby("District")["Packs_Given"].sum()
    s = pd.DataFrame({
        "Dry Food - Packs (incl. combo)": g(food),
        "Clothes - Packs (incl. combo)": g(rec["Item_Type"] == "Dry Food + Clothes"),
        "Books - Packs": g(rec["Item_Type"] == "Books"),
        "Total Packs": rec.groupby("District")["Packs_Given"].sum(),
        "Total Families": rec.groupby("District")["Affected_Families"].sum(),
        "Disaster-Mode Packs": g(rec["Mode"] == "Disaster"),
        "Normal-Mode Packs": g(rec["Mode"] == "Normal")}).fillna(0).astype(int)
    s.loc["TOTAL"] = s.sum()
    return s.reset_index(names="District")


def replace_sheet(wb, name, df):
    pos = wb.sheetnames.index(name)
    del wb[name]
    ws = wb.create_sheet(name, pos)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))


def main():
    if not BACKUP.exists():
        shutil.copy2(DIST_FILE, BACKUP)
        print(f"Original workbook backed up to {BACKUP.name}")
    rec = generate()
    wb = openpyxl.load_workbook(DIST_FILE)
    replace_sheet(wb, "Distribution_Records", rec)
    replace_sheet(wb, "Summary_By_District", district_summary(rec))

    readme = wb["README"]
    start = next(r for r in range(1, readme.max_row + 1)
                 if str(readme.cell(r, 1).value).startswith("DUMMY data"))
    for r in range(start, max(readme.max_row, start + len(README_DUMMY)) + 1):
        i = r - start
        readme.cell(r, 1).value = README_DUMMY[i] if i < len(README_DUMMY) else None
    wb.save(DIST_FILE)
    print(f"Wrote {len(rec)} simulated records ({(rec['Mode'] == 'Disaster').sum()} disaster mode) "
          f"across {rec['DS_Division'].nunique()} DS divisions")


if __name__ == "__main__":
    main()
